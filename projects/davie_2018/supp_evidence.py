#!/usr/bin/env python
"""Davie 2018: join the GEO cluster labels to the paper's Tables S2/S3.

Project-specific (Table S2/S3 layouts are this paper's), so it lives in the
project folder rather than scripts/. Reads
  - the GEO 57K metadata, grouped by the authors' per-cell label (`res.2` = dominant
    Seurat res-2.0 id)
  - supp/mmc2.xlsx              (Table S2: res-2.0 / res-8.0 / subcluster annotations)
  - supp/mmc3.xlsx              (Table S3: Seurat res-2.0 marker genes)
and writes
  - supp_evidence.json    per label: res-2.0 id, S2 rows for that id, top S3 markers
                          (the report agents' structured evidence)
The per-cluster Table S2 comparison now lives in build_clusters.py
(s2_annotation / author_conflict on each cluster). SYNONYMS / _norm / read_s2 are
shared with it.

Usage:
    uv run python projects/davie_2018/supp_evidence.py
"""

from __future__ import annotations

import csv
import gzip
import importlib.util
import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

PROJ = Path(__file__).resolve().parent
SOURCE_FILE = "GSE107451_DGRP-551_w1118_WholeBrain_57k_Metadata.tsv.gz"
N_MARKERS = 20

# Curated naming equivalences between GEO labels and Table S2 annotation strings
# (same cell set, different spelling / level of naming). Anything not covered here
# or by exact (normalised) match is flagged REVIEW.
SYNONYMS = {
    "A/B*-KC": "a'/b'-KC",                       # '*' = prime
    "Clock": "Clock neurons",
    "dorsal_Fan-shaped_Body": "dFB",
    "Plasmatocytes": "Hemocytes",                # plasmatocytes are the hemocyte type
    "Lamina_monopolar": "L1/L2/L3/L4/L5",        # unresolved remainder of the lamina cluster
    "PAM": "PAM (Fer2(+)) neurons",
    "Olfactory_projection_neurons": "OPN: adPN and lPN",
    "adPN": "OPN: adPN and lPN",
    "lPN": "OPN: adPN and lPN",
    "adPN/C15": "C15(+) adPNs",
    "adPN/C15&kn": "C15(+) kn(+) adPNs",
    "adPN/kn": "kn(+) adPNs",
    "lPN/unpg": "unpg(+) lPNs",
    "lPN/CG31676": "CG31676(+) lPNs",
    "Capa": "Capa neurons",
    "Gr43a": "Gr43a neurons",
    "Hug": "Hugin neurons",
}
# Labels from the paper's neuropeptide re-clustering ("all cells expressing
# neuropeptides", cells drawn from many res-2.0 clusters): a res-2.0 mismatch is
# expected for these, so they are not flagged on that basis.
NEUROPEPTIDE_SUBCLUSTERS = {
    "AstA/Nplp1", "AstA/NPF", "CCHa1", "Crz", "FMRFa", "ITP", "Mip/ITP", "Mip/OCT",
    "Proc/Ms", "Proc/Gpb5", "Capa", "Hug", "IPC", "CCAP", "LNv", "Gr43a",
}


def _norm(s: str) -> str:
    """Loose comparison key: case/space/underscore/prime-marker insensitive."""
    s = s.lower().replace("*", "'").replace("’", "'")
    return re.sub(r"[\s_\-]+", "", s)


def read_s2() -> tuple[dict[str, dict], list[dict]]:
    """Return ({res2_id: row}, [res8/subcluster rows]) from Table S2."""
    ws = openpyxl.load_workbook(PROJ / "supp" / "mmc2.xlsx", read_only=True).worksheets[0]
    res2: dict[str, dict] = {}
    other: list[dict] = []
    section = None
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in row[:5]]
        head = cells[0].replace("\n", " ")
        if head.startswith("Resolution 2.0"):
            section = "res2"; continue
        if head.startswith("Resolution 8.0"):
            section = "res8"; continue
        if head.startswith("Sub-clustering"):
            section = "sub"; continue
        if not any(cells) or head.startswith(("Cluster ID", "Input")):
            continue
        if section == "res2":
            res2[head] = {"cluster_id": head, "annotation": cells[1],
                          "method": cells[2], "references": cells[3]}
        elif section == "res8":
            other.append({"level": "res8", "cluster_id": head, "annotation": cells[1],
                          "method": cells[2], "references": cells[3]})
        elif section == "sub":
            other.append({"level": "subcluster", "input_res2_id": head,
                          "input_annotation": cells[1].replace("\n", " "),
                          "method": cells[2], "annotation": cells[3],
                          "references": cells[4]})
    return res2, other


def read_s3() -> dict[str, list[dict]]:
    ws = openpyxl.load_workbook(PROJ / "supp" / "mmc3.xlsx", read_only=True).worksheets[0]
    markers: dict[str, list[dict]] = defaultdict(list)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2 or row[2] is None:
            continue
        p, diff, cid, gene = row[:4]
        markers[str(cid)].append({"gene": str(gene), "avg_diff": float(diff), "p": float(p)})
    for cid in markers:
        markers[cid].sort(key=lambda m: (m["p"], -m["avg_diff"]))
    return markers


def label_annotations() -> list[dict]:
    """Per-GEO-label summary (cells, dominant res-2.0 id, timepoints), built straight
    from the GEO metadata with the generic label ingester (the label-level view the
    reports were written from; cell_type_annotations.json is now per cluster)."""
    spec = importlib.util.spec_from_file_location(
        "build_annotations_from_geo", PROJ.parents[1] / "scripts" / "build_annotations_from_geo.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    with gzip.open(PROJ.parents[1] / ".cache" / SOURCE_FILE, "rt") as fh:
        rows = list(csv.DictReader(fh, delimiter="\t"))
    return mod.build(rows, label_col="annotation", carry_cols=["res.2"], time_col="Age")


def main() -> None:
    anns = label_annotations()
    res2, other = read_s2()
    markers = read_s3()

    evidence = {}
    for a in anns:
        label, rid = a["label"], a.get("res.2", "")
        if label.isdigit():
            continue  # unannotated numbered clusters: held, no report needed
        s2 = res2.get(rid, {})
        related = [o for o in other
                   if _norm(o["annotation"]) == _norm(label)
                   or (o.get("input_res2_id") == rid)]
        evidence[label] = {
            "label": label,
            "cell_number": a["cell_number"],
            "timepoints_days": a.get("timepoints", []),
            "res2_cluster_id": rid,
            "res2_ambiguous": a.get("res.2_ambiguous"),
            "table_s2_res2": s2 or None,
            "table_s2_related": related,
            "table_s3_top_markers_for_res2_cluster": [
                m["gene"] for m in markers.get(rid, [])[:N_MARKERS]],
        }
    (PROJ / "supp_evidence.json").write_text(json.dumps(evidence, indent=2) + "\n")
    print(f"evidence for {len(evidence)} labels")


if __name__ == "__main__":
    main()
