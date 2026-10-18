#!/usr/bin/env python
"""Allen et al. 2020: cell_type_annotations.json from the re-run clusters and the
authors' per-cluster table.

Each cell's cluster comes from re-running the authors' pipeline (rerun/: all 120 cluster
sizes equal the paper's), written to .cache/allen/rerun/clusters.tsv.gz. The authors type
each cluster in the paper's Figure 1-source data 2 (supp/elife-54074-fig1-data2.xlsx):
neuronal clusters by neurotransmitter, predicted hemilineage, enriched peptides and Hox
genes; non-neuronal clusters by cell type. One entry per cluster:

  label              vnc_<n>
  type               the most specific author call: the predicted hemilineage (e.g.
                     "hemilineage 12B", "hemilineages 13A/19A"); else the neurotransmitter
                     class ("cholinergic neuron", ...); else the non-neuronal cell type;
                     else "neuron"
  author_*           the table's fields (neurotransmitter, hemilineage, peptides, Hox, cell type)
  unannotated        clusters the authors left "undefined" or typed only as neuron
                     (held: untyped)
  cells_by_sample    cells per Sample (sex x replicate)

Also writes .cache/allen/memberships.tsv.gz (cell<TAB>cluster) for the expression build.

Usage:
    uv run --with openpyxl python projects/allen_2020/build_clusters.py
"""

from __future__ import annotations

import csv
import gzip
import json
from collections import Counter
from pathlib import Path

import openpyxl
import yaml

PROJ = Path(__file__).resolve().parent
CACHE = PROJ.parents[1] / ".cache" / "allen"
TABLE = PROJ / "supp" / "elife-54074-fig1-data2.xlsx"
NT_CLASS = {"acetylcholine": "cholinergic neuron", "GABA": "GABAergic neuron",
            "glutamate": "glutamatergic neuron", "Monoaminergic": "monoaminergic neuron",
            "Peptidergic": "peptidergic neuron", "mixed": "neuron of mixed neurotransmitter identity"}
SAMPLE = {"Female_rep1": "female_rep_1", "Female_rep2": "female_rep_2",
          "Male_rep1": "male_rep_1", "Male_rep2": "male_rep_2"}


def table() -> dict[str, dict]:
    wb = openpyxl.load_workbook(TABLE, read_only=True)
    out = {}
    for row in list(wb["Neuronal Clusters"].iter_rows(values_only=True))[1:]:
        if row[0] is None:
            continue
        n, cells, _, _, nt, hl, pep, hox = row[:8]
        out[str(int(n))] = {"cells": int(cells), "neurotransmitter": nt, "hemilineage": hl,
                            "peptides": pep, "hox": hox, "class": "neuronal"}
    for row in list(wb["Non-Neuronal Clusters"].iter_rows(values_only=True))[1:]:
        if row[0] is None:
            continue
        out[str(int(row[0]))] = {"cells": int(row[1]), "cell_type": row[4], "class": "non-neuronal"}
    return out


def type_of(t: dict) -> str:
    if t["class"] == "non-neuronal":
        return t["cell_type"]
    hl = t.get("hemilineage")
    if hl:
        return f"hemilineages {hl}" if "/" in hl else f"hemilineage {hl}"
    return NT_CLASS.get(t.get("neurotransmitter"), "neuron")


def main() -> None:
    cfg = yaml.safe_load((PROJ / "config.yaml").read_text())
    (clustering,) = cfg["sample_build"]["clusterings"]
    tab = table()
    with gzip.open(CACHE / "rerun" / "clusters.tsv.gz", "rt") as fh:
        rows = list(csv.DictReader(fh, delimiter="\t"))
    by: dict[str, list[str]] = {}
    for r in rows:
        by.setdefault(r["cluster"], []).append(r["cell"])
    if set(by) != set(tab):
        raise SystemExit(f"clusters differ from the authors' table: {sorted(set(by) ^ set(tab))}")
    anns = []
    for n in sorted(by, key=int):
        t, cells = tab[n], by[n]
        if t["cells"] != len(cells):
            raise SystemExit(f"cluster {n}: {len(cells)} cells, {t['cells']} in the authors' table")
        rec = {"label": f"vnc_{n}", "type": type_of(t), "clustering": clustering["id"],
               "clustering_name": clustering["name"], "source_cluster_id": n, "cell_number": len(cells),
               "cells_by_sample": dict(sorted(Counter(SAMPLE[c.rsplit("_", 1)[0]] for c in cells).items())),
               "author_class": t["class"]}
        for k in ("neurotransmitter", "hemilineage", "peptides", "hox", "cell_type"):
            if t.get(k):
                rec[f"author_{k}"] = t[k]
        if rec["type"] in ("undefined", "neuron"):
            rec["unannotated"] = True
        anns.append(rec)
    with gzip.open(CACHE / "memberships.tsv.gz", "wt") as fh:
        fh.write("barcode\tcluster\n")
        fh.writelines(f"{r['cell']}\tvnc_{r['cluster']}\n" for r in rows)
    out = {
        "source": {"doi": cfg["source"]["doi"], "title": cfg["source"]["title"]},
        "dataset": {
            "accession": cfg["dataset"]["accession"],
            "label_source": "the authors' 120 clusters, cells assigned by re-running their published "
                            "pipeline (all cluster sizes equal the paper's); typed from Figure 1-source data 2",
            "label_level": "author cluster",
            "validate_against": "paper, Figure 1-source data 2, Figure 4",
        },
        "annotations": anns,
    }
    (PROJ / "cell_type_annotations.json").write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n")
    held = [a["label"] for a in anns if a.get("unannotated")]
    print(f"wrote {len(anns)} clusters, {len(rows)} cells; types: "
          f"{dict(Counter(a['type'] for a in anns).most_common())}; held: {held}")


if __name__ == "__main__":
    main()
